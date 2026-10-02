"""
Billback Voucher filler — SELF-CONTAINED, does not affect any other page.

Takes a bill-back 'Monthly List of Discount Participation' PDF, sums DA spend
per voucher brand code (using a voucher-only mapping with its own RE/REE,
CDN/CDA, OT-combined and QB-combined splits), and fills the MMD Billback
Voucher template's Depletions (4005) column. Distributor/State is read from
the PDF; all other header fields are left blank for manual entry.
"""
import io
import os
import re
from datetime import date

import engine  # reuse the proven, market-agnostic parser (read-only)


# ---- Voucher-only line -> code classifier ----
# Returns the voucher code for a product description, applying the splits that
# exist ONLY for the voucher (these do not touch brand_hierarchy.py).
# Order matters: most specific first.
def _voucher_code(desc):
    d = (desc or "").upper()

    # --- Roederer family ---
    if "CRISTAL" in d or "BRUT CRIS" in d:
        return "LRC"                      # LRC = Cristal (incl. "BRUT CRIS" gift packs)
    if "ROEDERER EST" in d or "L'ERMITAGE" in d or "LERMITAGE" in d:
        # Estate split: L'Ermitage -> REE, everything else Estate -> RE
        if "ERMITAGE" in d:
            return "REE"
        return "RE"
    if "DEMI SEC" in d:
        return "LRV"                      # Demi Sec is a VRBB line
    if "BRUT COL" in d or "CARTE BLANCHE" in d or "BRUT PREMIER" in d or "COLLECTION" in d:
        return "LR"                       # LR = Collection
    # plain Louis Roederer Brut <vintage>, Blanc de Blanc, Brut Rose/Nature -> VRBB
    if "LOUIS ROEDERER" in d and (
        re.search(r"BRUT \d{2}", d) or "BLANC DE BL" in d or "BL DE BL" in d
        or "BRUT ROSE" in d or "BRUT NATURE" in d or "BRUT NAT" in d or "BRUT VINTAGE" in d
    ):
        return "LRV"

    # --- Ott (combined) ---
    if "DOM OTT" in d or "BY OTT" in d:
        return "OT"

    # --- Querciabella (combined incl. Mongrana) ---
    if "QUERCIABELLA" in d:
        return "QB"

    # --- Carpe Diem split ---
    if "CARPE DIEM" in d:
        if "ANDERSON" in d:
            return "CDA"                  # Anderson Valley
        return "CDN"                      # everything else Carpe Diem -> Napa

    # --- Dominus family ---
    if "NAPANOOK" in d:
        return "DEN"
    if "OTHELLO" in d:
        return "DEO"
    if "DOMINUS" in d:
        return "DE"

    # --- single-code brands ---
    if "SCHARFFENBERGER" in d:
        return "SC"
    if "DOMAINE ANDERSON" in d:
        return "DA"
    if "MERRY EDWARDS" in d:
        return "MEW/SVC"
    if "INNISKILLIN" in d:
        return "AW"                       # Aterra Wines = Inniskillin
    if "JACKSON TRIGGS" in d:
        return "AWT"
    if "BON VIVANT" in d:
        return "BV"
    if "PICHON" in d:
        return "PL"
    if "CH DE PEZ" in d or "BORDEAUX" in d and "CLR" in d:
        return "BX"                       # CLR Bordeaux
    if any(k in d for k in ["CH DE SALES", "CH PEYMOUTON", "PEYMOUTON", "PUY BLANQUET",
                            "LAFLEUR-GAZIN", "BOURGNEUF", "CERTAN", "HOSANNA",
                            "MAGDELAINE", "PETRUS", "GIACONDA"]):
        return "MX"                       # Ets. J-P Moueix
    if "LOUDENNE" in d:
        return "LO"
    if "DELAS" in d:
        return "DF"                       # Delas Freres
    if "LADOUCETTE" in d or "MARC BREDIF" in d or "LA POUSSIE" in d or "COMTE LAFOND" in d \
       or "LES DEUX TOURS" in d or "BARON DE L" in d or "REGNARD" in d:
        return "LV"                       # LV = Loire Valley (Regnard placed here for now)
    if "SCHLUMBERGER" in d:
        return "DS"                       # Domaine Schlumberger
    if "PIO CESARE" in d:
        return "PC"
    if "ESPERTO" in d:
        return "LFE"                      # Esperto
    if "FELLUGA" in d:
        return "LF"                       # Livio Felluga
    if "CASTIGLION" in d:
        return "CB"                       # Castiglion del Bosco (CB / CS)
    if "RAMOS PINTO" in d:
        return "RP"
    if "DUAS QUINTAS" in d:
        return "DQ"
    if "MURRIETA" in d or "DALMAU" in d or "CAPELLANIA" in d or "PAZO BARRANTES" in d:
        return "MU"
    if "FLEUR DU CAP" in d:
        return "SA"                       # Distell
    if "MEERLUST" in d:
        return "ME"

    return None  # unmapped -> reported back to the user


# Voucher code -> row in the template (column B labels)
_CODE_ROW = {
    "LRC": 15, "LR": 16, "LRV": 17, "RE": 18, "REE": 19, "SC": 20,
    "DE": 21, "DEN": 22, "DEO": 23, "DA": 24, "CDN": 25, "CDA": 26,
    "MEW/SVC": 27, "AW": 28, "AWT": 29, "BV": 30, "PL": 31, "BX": 32,
    "MX": 33, "LO": 34, "DF": 35, "OT": 36, "LV": 37, "DS": 38, "QB": 39,
    "PC": 40, "LF": 41, "LFE": 42, "CB": 43, "RP": 44, "DQ": 45, "MU": 46,
    "SA": 47, "ME": 48,
}
_DEPLETIONS_COL = 3   # column C
_TOTAL_ROW = 49


def summarize_for_voucher(pdf_path):
    """Parse one bill-back PDF -> (state, period_end_date, code_totals, unmapped).

    code_totals: {voucher_code: summed DA spend}
    unmapped: list of (desc, ext) lines that didn't map to a code
    """
    lines, mdate, market, stated = engine.parse_participation_pdf(pdf_path)
    code_totals = {}
    unmapped = []
    for l in lines:
        code = _voucher_code(l["desc"])
        if code is None:
            unmapped.append((l["desc"], l["ext"]))
            continue
        code_totals[code] = code_totals.get(code, 0.0) + l["ext"]
    # Distributor/State string straight from the PDF header
    state = _read_state(pdf_path)
    return state, mdate, code_totals, unmapped


def _read_state(pdf_path):
    """Read the 'SGWS ...' distributor/state string from the PDF header."""
    import pdfplumber
    try:
        with pdfplumber.open(pdf_path) as pdf:
            full = "\n".join((p.extract_text() or "") for p in pdf.pages[:1])
    except Exception:
        return ""
    m = re.search(r"(SGWS[^\n]*?)(?:\s+Date\s*:|\n|$)", full)
    if m:
        # Trim trailing report-date fragments if any slipped in
        return re.sub(r"\s+Date\s*:.*$", "", m.group(1)).strip()
    return ""


def fill_voucher(pdf_path, template_path):
    """Return (xlsx_bytes, state, code_totals, unmapped) with the voucher filled.

    Fills only: Distributor/State (F1) and the Depletions column per brand code.
    Leaves Invoice #, Invoice date, Date, Name/Region, Signature, Approval blank.
    """
    import openpyxl
    state, mdate, code_totals, unmapped = summarize_for_voucher(pdf_path)

    wb = openpyxl.load_workbook(template_path)
    ws = wb["Sheet1"]

    # Distributor/State -> F1 (label "Distributor, State : " is in E1)
    if state:
        ws.cell(1, 6).value = state

    # Clear header fields that should be filled in manually (some are pre-filled
    # in the template). Leave blank: Invoice #, Invoice date, Date, Name/Region,
    # Signature, Approval.
    ws.cell(2, 6).value = None    # F2  Invoice #
    ws.cell(2, 10).value = None   # J2  Invoice date
    ws.cell(4, 10).value = None   # J4  Date
    ws.cell(5, 10).value = None   # J5  Name / Region
    ws.cell(6, 10).value = None   # J6  Signature
    ws.cell(7, 10).value = None   # J7  Approval

    # Depletions column per code
    total = 0.0
    for code, amt in code_totals.items():
        row = _CODE_ROW.get(code)
        if not row:
            continue
        ws.cell(row, _DEPLETIONS_COL).value = round(amt, 2)
        total += amt
    # Total row (Depletions column)
    ws.cell(_TOTAL_ROW, _DEPLETIONS_COL).value = round(total, 2)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue(), state, code_totals, unmapped
