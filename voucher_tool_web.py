"""
DEPLOY-READY Billback Voucher Tool  —  web version with Microsoft (Entra ID) login.

Same tool as voucher_tool.py, plus:
  - Microsoft 365 sign-in (OIDC via Entra ID), restricted to @mmdusa.com
  - Shared learned-mappings stored on a persistent path (team-wide, survives redeploys)
Your existing DA app files are NOT modified; this imports voucher.py unchanged.

Local run (no login):   py -m streamlit run voucher_tool_web.py
  -> if no auth secrets are configured, it runs open (for local testing).
Deployed run (with login): configure [auth] in .streamlit/secrets.toml (see README_DEPLOY.txt).
"""
import io, os, json, re
import streamlit as st
import openpyxl
import voucher as V

# ---- storage path: a persistent disk in prod, local file in dev ----
DATA_DIR = os.environ.get("VOUCHER_DATA_DIR", ".")
LEARNED_PATH = os.path.join(DATA_DIR, "learned_voucher_map.json")

CODES = list(V._CODE_ROW.keys())
DEPLETIONS_COL = V._DEPLETIONS_COL
CODE_ROW = V._CODE_ROW
TOTAL_ROW = V._TOTAL_ROW
ALLOWED_DOMAIN = os.environ.get("ALLOWED_EMAIL_DOMAIN", "mmdusa.com")

st.set_page_config(page_title="MMD Billback Voucher", page_icon="🧾", layout="wide")

# ============ AUTH (Microsoft Entra ID via Streamlit OIDC) ============
def require_login():
    """If auth is configured, enforce Microsoft login + domain restriction.
    If not configured (local dev), allow through so the tool still runs."""
    has_auth = False
    try:
        has_auth = bool(st.secrets.get("auth"))
    except Exception:
        has_auth = False
    if not has_auth:
        return True  # local/dev: no gate
    if not getattr(st.user, "is_logged_in", False):
        st.title("🧾 MMD Billback Voucher")
        st.write("Please sign in with your MMD Microsoft account.")
        st.button("Sign in with Microsoft", on_click=st.login)
        st.stop()
    email = (getattr(st.user, "email", "") or "").lower()
    if ALLOWED_DOMAIN and not email.endswith("@" + ALLOWED_DOMAIN):
        st.error(f"Access restricted to @{ALLOWED_DOMAIN} accounts. You are signed in as {email or 'unknown'}.")
        st.button("Sign out", on_click=st.logout)
        st.stop()
    return True

require_login()

# ============ learned mappings (shared, persistent) ============
def load_learned():
    try:
        with open(LEARNED_PATH) as f:
            return json.load(f)
    except Exception:
        return {}

def save_learned(m):
    try:
        os.makedirs(DATA_DIR, exist_ok=True)
        with open(LEARNED_PATH, "w") as f:
            json.dump(m, f, indent=2)
        return True
    except Exception:
        return False

# ============ the two baked-in fixes over the imported classifier ============
def classify(desc, learned):
    d = (desc or "").upper()
    if desc in learned:
        return learned[desc]
    if "CLOS JORDANNE" in d:
        return "AW"
    if "LOUIS ROEDERER" in d and re.search(r"BRUT \d{2}", d) \
       and not any(k in d for k in ("BRUT COL", "BRUT CRIS", "CRISTAL", "CARTE BLANCHE", "BRUT PREMIER")):
        return "LRV"
    return V._voucher_code(desc)

def parse_lines(pdf_path):
    import engine as ENG
    lines, mdate, market, stated = ENG.parse_participation_pdf(pdf_path)
    return lines, mdate, stated

def read_state(pdf_path):
    try:
        return V._read_state(pdf_path)
    except Exception:
        return ""

def build_workbook(code_totals, state):
    wb = openpyxl.load_workbook("voucher_template.xlsx")
    ws = wb["Sheet1"]
    if state:
        ws.cell(1, 6).value = state
    total = 0.0
    for code, amt in code_totals.items():
        row = CODE_ROW.get(code)
        if not row:
            continue
        ws.cell(row, DEPLETIONS_COL).value = round(amt, 2)
        total += amt
    ws.cell(TOTAL_ROW, DEPLETIONS_COL).value = round(total, 2)
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf.getvalue(), total

# ============ UI ============
st.title("🧾 MMD Billback Voucher")
cap = "Fill the MMD voucher from a bill-back PDF; resolve any exceptions with a dropdown."
try:
    if st.secrets.get("auth") and getattr(st.user, "is_logged_in", False):
        cap += f"  ·  Signed in as {st.user.email}"
except Exception:
    pass
st.caption(cap)

learned = load_learned()

with st.sidebar:
    st.header("Input")
    pdf_file = st.file_uploader("Bill-back PDF", type=["pdf"])
    st.markdown("---")
    st.caption(f"{len(learned)} shared learned mapping(s)")
    try:
        if st.secrets.get("auth") and getattr(st.user, "is_logged_in", False):
            st.button("Sign out", on_click=st.logout)
    except Exception:
        pass

if not pdf_file:
    st.info("Upload a bill-back PDF to begin.")
    st.stop()

tf = os.path.join(DATA_DIR, "_uploaded_voucher.pdf")
with open(tf, "wb") as f:
    f.write(pdf_file.getbuffer())

try:
    lines, mdate, stated = parse_lines(tf)
except Exception as e:
    import traceback
    st.error(f"Parse failed: {e}"); st.code(traceback.format_exc()); st.stop()

state = read_state(tf)

mapped = {}; exceptions = []; pdf_total = 0.0
for l in lines:
    pdf_total += l["ext"]
    code = classify(l["desc"], learned)
    if code == "__IGNORE__":
        continue
    if code is None:
        exceptions.append(l)
    else:
        mapped[code] = mapped.get(code, 0.0) + l["ext"]

exc_by_desc = {}
for e in exceptions:
    exc_by_desc[e["desc"]] = exc_by_desc.get(e["desc"], 0.0) + e["ext"]

c1, c2, c3 = st.columns(3)
c1.metric("PDF total", f"${pdf_total:,.2f}")
c1.caption(f"Distributor/State: {state or '(not detected)'}  ·  {mdate or ''}")
c2.metric("Mapped", f"${sum(mapped.values()):,.2f}")
exc_total = sum(exc_by_desc.values())
c3.metric("Unassigned", f"${exc_total:,.2f}",
          delta=None if exc_total == 0 else f"{len(exc_by_desc)} line(s)", delta_color="inverse")

assignments = {}
if exc_by_desc:
    st.subheader("⚠ Exceptions — assign a code to each")
    st.caption("Pick a code (or Ignore). Saved choices are shared with the team and remembered next time. "
               "You can download without assigning — unassigned lines just won't be on the voucher yet.")
    for i, (desc, amt) in enumerate(sorted(exc_by_desc.items(), key=lambda x: -x[1])):
        col1, col2 = st.columns([3, 1])
        col1.markdown(f"**{desc}**  \n${amt:,.2f}")
        assignments[desc] = col2.selectbox("Code", ["— unassigned —", "Ignore this line"] + CODES,
                                           key=f"exc_{i}", label_visibility="collapsed")
else:
    st.success("No exceptions — every line classified cleanly.")

final = dict(mapped); still = 0.0; new_learned = {}
for desc, amt in exc_by_desc.items():
    ch = assignments.get(desc, "— unassigned —")
    if ch == "Ignore this line":
        new_learned[desc] = "__IGNORE__"
    elif ch in CODES:
        final[ch] = final.get(ch, 0.0) + amt; new_learned[desc] = ch
    else:
        still += amt

st.subheader("Voucher preview (Depletions)")
prev = sorted(final.items(), key=lambda x: -x[1])
st.dataframe({"Code": [c for c, _ in prev], "Depletions $": [f"${a:,.2f}" for _, a in prev]})
built = sum(final.values())
if abs(built - pdf_total) < 0.01:
    st.success(f"Voucher total ${built:,.2f} ties to the PDF exactly.")
else:
    st.warning(f"Voucher ${built:,.2f} vs PDF ${pdf_total:,.2f} — ${still:,.2f} still unassigned.")

cA, cB = st.columns(2)
if cA.button("💾 Save code choices (shared with team)", disabled=not new_learned):
    if save_learned({**learned, **new_learned}):
        st.success(f"Saved {len(new_learned)} mapping(s)."); st.rerun()
    else:
        st.error("Could not write learned map (check storage path/permissions).")

data, _ = build_workbook(final, state)
cB.download_button("⬇ Download voucher (Excel)", data=data,
    file_name="Billback_Voucher_filled.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
