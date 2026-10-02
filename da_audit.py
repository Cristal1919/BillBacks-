"""
DA Reconciliation Audit Engine
==============================
For MMD USA — California (NCA + SCA) Discount Allowance reconciliation.

Methodology (locked in after 4 iterations):
  1. One row per BILL-BACK LINE (BB is the authoritative DA invoice).
  2. VIP cases at same item × region × price band give the premise mix (ON/OFF).
  3. For each premise: find program tier with CLOSEST PTA. That tier authorizes a DA.
  4. Per-case overage = max(0, BB DA − authorized DA). Negative differences are $0
     (we don't ask SGWS for more when they undercharged).
  5. Total line overage = ON overage + OFF overage.
  6. Different MMD IDs are different products. Vintage variants collapse to base.
  7. Special handling: an MMD ID can be routed to a SPECIFIC TIER COLUMN inside
     another product's program block (e.g. RE1110028 → RE1100028 'Special cuvee' tier).

USAGE: see audit_runner.py — orchestrates the full audit and writes the workbook.
"""

import re
import pickle
from collections import defaultdict
from openpyxl import load_workbook


# Standard tier names across MMD program files
ON_TIERS = ['WTC', 'Glass Pour', 'Glass Pour Hot', 'Glass Pour Hotel',
            'BTG', 'Large BTG', 'New BTG', 'Ambassador', 'Netjets', 'NetJets',
            'Special cuvee']
OFF_TIERS = ['Front Line', '1 Cs Deal', '2 Cs Deal', '3 Cs Deal', '5 Cs Deal',
             '10 Cs Deal', 'Chain', 'Key retailers', 'Key Retailers',
             'Costco', 'Special Feature']


# ============================================================
# PROGRAM FILE PARSER
# ============================================================
def parse_program_file(program_path, brand_prefix='RE'):
    """
    Parse an MMD program Excel file (DI / DI 10% / NDI sheets).
    
    Returns dict with:
      - all_blocks: list of program tier blocks
      - mmd_to_blocks: {mmd_id: [block, ...]}
      - workbook: openpyxl workbook (for tier lookups later)
    
    Each block has: sheet, header_row, pta_row, da_row, fob, product,
                    mmd_ids (list), tiers ({col: tier_name})
    """
    wb = load_workbook(program_path, data_only=True)
    all_blocks = []
    
    for sheet_name in wb.sheetnames:
        if 'Freight' in sheet_name or 'Tax' in sheet_name:
            continue
        ws = wb[sheet_name]
        
        # Find blocks: each starts with a row containing 'Front Line' in col 2
        starts = []
        for r in range(1, ws.max_row + 1):
            v = ws.cell(r, 2).value
            if v and str(v).strip() == 'Front Line':
                starts.append(r)
        
        for i, hr in enumerate(starts):
            next_hr = starts[i + 1] if i + 1 < len(starts) else ws.max_row + 1
            block = {
                'sheet': sheet_name,
                'header_row': hr,
                'block_label': str(ws.cell(hr, 1).value or '').strip(),
                'tiers': {},
                'fob': None,
                'mmd_ids': [],
                'product': None,
                'fob_row': None,
                'pta_row': None,
                'da_row': None,
            }
            # Capture tier columns (column 2 onwards)
            for c in range(2, 17):
                h = ws.cell(hr, c).value
                if h and str(h).strip() != 'Blank':
                    block['tiers'][c] = str(h).strip()
            
            # Scan rows after header for FOB, PTA, DA, MMD IDs, product name
            for off in range(1, min(next_hr - hr, 18)):
                rr = hr + off
                if rr > ws.max_row:
                    break
                lbl = str(ws.cell(rr, 1).value or '').strip().lower()
                if lbl.startswith('fob') and not block['fob']:
                    block['fob'] = ws.cell(rr, 2).value
                    block['fob_row'] = rr
                elif 'mmd da support' in lbl:
                    block['da_row'] = rr
                elif 'price to acct per case' in lbl:
                    block['pta_row'] = rr
                
                # MMD IDs typically in column 19
                v19 = ws.cell(rr, 19).value
                if v19 and brand_prefix in str(v19):
                    ids_str = str(v19).strip()
                    for piece in ids_str.replace(', ', ',').split(','):
                        piece = piece.strip()
                        if piece and piece not in block['mmd_ids']:
                            block['mmd_ids'].append(piece)
                
                if not block['product']:
                    v = ws.cell(rr, 17).value
                    if v:
                        block['product'] = str(v).strip()
            
            # Fallback PTA detection
            if not block['pta_row'] and block['da_row']:
                rr = block['da_row'] + 1
                if rr <= ws.max_row:
                    tier_vals = []
                    for col in block['tiers']:
                        v = ws.cell(rr, col).value
                        if isinstance(v, (int, float)):
                            tier_vals.append(v)
                    if tier_vals and max(tier_vals) > 50:
                        block['pta_row'] = rr
            
            if block['fob']:
                all_blocks.append(block)
    
    mmd_to_blocks = defaultdict(list)
    for b in all_blocks:
        for mmd in b['mmd_ids']:
            mmd_to_blocks[mmd].append(b)
    
    return {
        'all_blocks': all_blocks,
        'mmd_to_blocks': dict(mmd_to_blocks),
        'workbook': wb,
        'path': program_path,
    }


def get_program_tiers(program, mmd_id, premise, item_routing=None):
    """
    Get all program tiers for an item at a given premise.
    
    item_routing (optional): {mmd_id: {'base_mmd': X, 'restrict_tiers': [tier_names] or None}}
      Lets you route an item to a different product's tier list, optionally
      restricted to specific tier columns (e.g. RE1110028 → RE1100028 'Special cuvee' only).
    """
    routing = (item_routing or {}).get(mmd_id, {'base_mmd': mmd_id, 'restrict_tiers': None})
    base_mmd = routing['base_mmd']
    restrict = routing.get('restrict_tiers')
    
    mmd_to_blocks = program['mmd_to_blocks']
    wb = program['workbook']
    
    if base_mmd not in mmd_to_blocks:
        return []
    
    target_tiers = ON_TIERS if premise == 'ON' else OFF_TIERS
    out = []
    for block in mmd_to_blocks[base_mmd]:
        if not block['pta_row'] or not block['da_row']:
            continue
        ws = wb[block['sheet']]
        for col, tier in block['tiers'].items():
            if tier not in target_tiers:
                continue
            if restrict and tier not in restrict:
                continue
            pta = ws.cell(block['pta_row'], col).value
            da = ws.cell(block['da_row'], col).value
            if not isinstance(pta, (int, float)) or pta == 0:
                continue
            if not isinstance(da, (int, float)):
                continue
            out.append({
                'tier': tier, 'pta': pta, 'da': da,
                'sheet': block['sheet'], 'col': col
            })
    return out


# ============================================================
# VIP DEPLETION PARSER
# ============================================================
def parse_vip_file(vip_path, brand_prefix='RE'):
    """
    Parse VIP depletion file (Pricing_Analysis_NCA___SCA.xlsx format).
    Auto-detects column layout (with or without 'Brands' column).
    
    Returns list of dicts: market, premise, item, name, size, price, cases.
    Filtered to brand_prefix and to ON/OFF premise.
    """
    wb = load_workbook(vip_path, data_only=True)
    ws = wb.active
    
    # Find header row: look for 'Markets' or 'Item Name ID' in any column of first 15 rows
    header_row = None
    col_map = {}
    for r in range(1, min(15, ws.max_row + 1)):
        for c in range(1, min(15, ws.max_column + 1)):
            v = str(ws.cell(r, c).value or '').strip()
            if v in ('Markets', 'OnOff Premises', 'Item Names', 'Item Name ID',
                     'Package Sizes', 'Price', 'Physical Cases'):
                if header_row is None:
                    header_row = r
                col_map[v] = c
    
    if header_row is None or 'Item Name ID' not in col_map:
        raise ValueError(f"Could not locate header row in {vip_path}")
    
    # Get column indices (file may or may not have a 'Brands' column)
    c_market = col_map.get('Markets', 1)
    c_premise = col_map.get('OnOff Premises', 2)
    c_name = col_map.get('Item Names', 3)
    c_iid = col_map['Item Name ID']
    c_size = col_map.get('Package Sizes', c_iid + 1)
    c_price = col_map.get('Price', c_iid + 2)
    c_cases = col_map.get('Physical Cases', c_iid + 3)
    
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        market = ws.cell(r, c_market).value
        premise = ws.cell(r, c_premise).value
        name = ws.cell(r, c_name).value
        iid = ws.cell(r, c_iid).value
        size = ws.cell(r, c_size).value
        price = ws.cell(r, c_price).value
        cases = ws.cell(r, c_cases).value
        
        if not iid or str(iid).strip() == 'Total':
            continue
        if not isinstance(cases, (int, float)) or cases == 0:
            continue
        if premise not in ('ON', 'OFF'):
            continue
        
        iid = str(iid).strip()
        if not iid.startswith(brand_prefix):
            continue
        
        market_short = 'NCA' if 'NCA' in str(market) else 'SCA' if 'SCA' in str(market) else '?'
        rows.append({
            'market': market_short, 'premise': premise, 'name': name,
            'item': iid, 'size': size, 'price': price, 'cases': cases,
        })
    
    return rows


# ============================================================
# BILL-BACK PDF PARSER
# ============================================================
def parse_billback_pdf(pdf_path, region, brand_markers=None, brand_prefix=None,
                       material_to_mmd=None):
    """
    Extract bill-back lines from SGWS bill-back PDF.
    
    Two ways to filter:
      - brand_markers: list of strings to require in description (e.g. ['DOM OTT'])
      - material_to_mmd + brand_prefix: filter by mapped MMD ID prefix
    
    If material_to_mmd is provided, that takes precedence (more accurate).
    
    Returns list of {region, line, mat, scc, desc, size, band, cases, da, tot}.
    """
    try:
        import pdfplumber
    except ImportError:
        raise ImportError("Install pdfplumber: pip install pdfplumber")
    
    # Pattern matches the SGWS bill-back line format across all markets.
    # Two market-specific variations are handled:
    #   - Product-group field (4 digits after material #) is OPTIONAL — Texas
    #     omits it; CA/HI/IL/NV/WA include it.
    #   - The "Less than" threshold may carry a comma (values >= 1,000) and may
    #     have no space after "than" (e.g. "Less than1,890.00"). Allowing
    #     optional space + commas captures the high-value lines (Cristal,
    #     Dominus, Brunello) that would otherwise silently drop out.
    pattern = re.compile(
        r'^(\d{7})\s+(?:(\d{4})\s+)?\d+\(WINE\)\s+\d+\s+(.+?)\s+'
        r'(\d+\.?\d*)\s+Less than\s*([\d,]+\.?\d*)\s+'
        r'(-?[\d,]+\.?\d*-?)\s+(-?[\d,]+\.?\d*)\s+(-?[\d,]+\.?\d*-?)'
    )
    
    out = []
    line_no = 0
    
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ''
            for raw in text.split('\n'):
                m = pattern.match(raw)
                if not m:
                    continue
                
                mat = m.group(1)
                
                # Filter logic: prefer material→MMD mapping; fall back to brand_markers
                if material_to_mmd and brand_prefix:
                    mmd = material_to_mmd.get(mat)
                    if not mmd or not str(mmd).startswith(brand_prefix):
                        continue
                elif brand_markers:
                    if not any(mk in raw.upper() for mk in brand_markers):
                        continue
                else:
                    pass  # no filter — return everything
                
                scc = m.group(2)
                desc_full = m.group(3).strip()
                size_m = re.search(r'(\d+(?:\.\d+)?)(ML|LT)', desc_full)
                size = size_m.group(0) if size_m else ''
                desc = re.sub(r'\d+(?:\.\d+)?(ML|LT)\s*$', '', desc_full).strip()
                
                net_price = float(m.group(4))
                band = float(m.group(5).replace(',', ''))
                cs = m.group(6).replace(',', '')
                da = m.group(7).replace(',', '')
                tot = m.group(8).replace(',', '')
                
                cases = -float(cs[:-1]) if cs.endswith('-') else float(cs)
                da_val = -float(da[:-1]) if da.endswith('-') else float(da)
                tot_val = -float(tot[:-1]) if tot.endswith('-') else float(tot)
                
                line_no += 1
                out.append({
                    'region': region, 'line': line_no,
                    'mat': mat, 'scc': scc, 'desc': desc, 'size': size,
                    'net_price': net_price, 'band': band,
                    'cases': cases, 'da': da_val, 'tot': tot_val,
                })
    
    return out


# ============================================================
# MATERIAL → MMD MAPPING (from CA distributor deal menu)
# ============================================================
def build_material_to_mmd_mapping(ca_deal_menu_path, brand_prefix='RE'):
    """Build {material_code: mmd_id} from CA distributor pricing file."""
    wb = load_workbook(ca_deal_menu_path, data_only=True)
    ws = wb.active
    
    # Headers at row 1; Item # in col 6, MMD ID in col 7
    mapping = {}
    for r in range(2, ws.max_row + 1):
        item = ws.cell(r, 6).value
        mmd = ws.cell(r, 7).value
        if not item or not mmd:
            continue
        if not str(mmd).startswith(brand_prefix):
            continue
        item_padded = str(item).strip().zfill(7)
        mapping[item_padded] = str(mmd).strip()
    
    return mapping


def extract_ca_deals(ca_deal_menu_path, brand_prefix='RE'):
    """Extract active deals from CA distributor menu for the brand."""
    wb = load_workbook(ca_deal_menu_path, data_only=True)
    ws = wb.active
    
    deals = []
    for r in range(2, ws.max_row + 1):
        mmd = ws.cell(r, 7).value
        if not mmd or not str(mmd).startswith(brand_prefix):
            continue
        end = ws.cell(r, 21).value
        end_str = str(end)[:10] if end else ''
        is_active = '9999' in end_str or '2026' in end_str or '2027' in end_str
        
        deals.append({
            'mmd_id': str(mmd).strip(),
            'desc': str(ws.cell(r, 8).value or ''),
            'size': ws.cell(r, 9).value,
            'discount_program': str(ws.cell(r, 11).value or ''),
            'net_case_price': ws.cell(r, 18).value,
            'fob': ws.cell(r, 24).value,
            'spa': ws.cell(r, 25).value or 0,
            'discount_support': ws.cell(r, 33).value or 0,
            'adjusted_laid_in': ws.cell(r, 37).value,
            'is_active': is_active,
            'premise': 'ON' if 'OSO' in str(ws.cell(r, 11).value or '').upper()
                       else ('OFF' if 'OFF' in str(ws.cell(r, 11).value or '').upper() else 'NONE'),
        })
    
    return deals


# ============================================================
# AUDIT ENGINE
# ============================================================
def run_audit(billback_lines, vip_rows, program, material_to_mmd,
              item_routing=None, default_premise='OFF'):
    """
    Run the per-bill-back-line audit.
    
    Returns list of audit records, one per BB line.
    Each record: bb, mmd, base_mmd, cases_by_premise, authorized_da_by_premise,
                 matched_tier_by_premise, overage_by_premise, compliant_dollars_by_premise,
                 overage_total, vip_matches, total_vip_cases, no_vip_match, not_in_program
    """
    item_routing = item_routing or {}
    mmd_to_blocks = program['mmd_to_blocks']
    
    # Build VIP lookup: 
    #   - If item routes to a base WITHOUT tier restriction → key by base (vintage collapse)
    #   - If item routes WITH tier restriction → key by original item (sub-product)
    #   - Otherwise → key by item itself
    def vip_key_for(item):
        routing = item_routing.get(item)
        if routing and not routing.get('restrict_tiers'):
            return routing['base_mmd']
        return item
    
    vip_by_key = defaultdict(list)
    for v in vip_rows:
        key_item = vip_key_for(v['item'])
        vip_by_key[(v['market'], key_item, v['price'])].append(v)
    
    audit = []
    for bb in billback_lines:
        mmd = material_to_mmd.get(bb['mat'])
        routing = item_routing.get(mmd, {'base_mmd': mmd})
        base_mmd = routing['base_mmd']
        vip_lookup_item = mmd if routing.get('restrict_tiers') else base_mmd
        
        rec = {
            'bb': bb, 'mmd': mmd, 'base_mmd': base_mmd,
            'cases_by_premise': {'ON': 0, 'OFF': 0},
            'authorized_da_by_premise': {'ON': 0, 'OFF': 0},
            'matched_tier_by_premise': {'ON': None, 'OFF': None},
            'overage_by_premise': {'ON': 0, 'OFF': 0},
            'compliant_dollars_by_premise': {'ON': 0, 'OFF': 0},
            'vip_matches': [], 'total_vip_cases': 0,
            'no_vip_match': True, 'overage_total': 0,
            'not_in_program': False,
        }
        
        # Skip credits/returns (negative cases or amounts)
        if bb['cases'] <= 0 or bb['tot'] <= 0:
            rec['note'] = 'Credit/return — skipped from overpayment calc'
            audit.append(rec)
            continue
        
        if not mmd:
            rec['note'] = 'No MMD mapping found for material'
            rec['not_in_program'] = True
            rec['cases_by_premise'][default_premise] = bb['cases']
            rec['overage_by_premise'][default_premise] = bb['tot']
            rec['overage_total'] = bb['tot']
            audit.append(rec)
            continue
        
        # Item not in program at all → entire amount unauthorized
        if base_mmd not in mmd_to_blocks:
            # Use VIP for premise split if available
            key = (bb['region'], vip_lookup_item, bb['band'])
            matches = vip_by_key.get(key, [])
            total_vip = sum(v['cases'] for v in matches)
            for prem in ('ON', 'OFF'):
                prem_cs = sum(v['cases'] for v in matches if v['premise'] == prem)
                if total_vip > 0:
                    alloc = bb['cases'] * (prem_cs / total_vip)
                elif prem_cs > 0:
                    alloc = bb['cases']
                else:
                    alloc = bb['cases'] if prem == default_premise else 0
                rec['cases_by_premise'][prem] = alloc
                rec['overage_by_premise'][prem] = alloc * bb['da']
            rec['vip_matches'] = matches
            rec['total_vip_cases'] = total_vip
            rec['no_vip_match'] = total_vip == 0
            rec['not_in_program'] = True
            rec['overage_total'] = sum(rec['overage_by_premise'].values())
            audit.append(rec)
            continue
        
        # Find matching VIP rows
        key = (bb['region'], vip_lookup_item, bb['band'])
        matches = vip_by_key.get(key, [])
        total_vip = sum(v['cases'] for v in matches)
        rec['vip_matches'] = matches
        rec['total_vip_cases'] = total_vip
        rec['no_vip_match'] = total_vip == 0
        
        # Closest PTA per premise
        for prem in ('ON', 'OFF'):
            tiers = get_program_tiers(program, mmd, prem, item_routing)
            if tiers:
                closest = min(tiers, key=lambda t: abs(t['pta'] - bb['band']))
                rec['matched_tier_by_premise'][prem] = closest
                rec['authorized_da_by_premise'][prem] = closest['da']
        
        # Allocate cases by VIP premise mix
        if total_vip > 0:
            for prem in ('ON', 'OFF'):
                prem_cs = sum(v['cases'] for v in matches if v['premise'] == prem)
                alloc = bb['cases'] * (prem_cs / total_vip)
                rec['cases_by_premise'][prem] = alloc
        else:
            # No VIP match — default to default_premise
            rec['cases_by_premise'][default_premise] = bb['cases']
        
        # Compute overage and compliant $ per premise
        for prem in ('ON', 'OFF'):
            cs = rec['cases_by_premise'][prem]
            auth_da = rec['authorized_da_by_premise'][prem]
            diff = bb['da'] - auth_da
            if diff > 0:
                rec['overage_by_premise'][prem] = cs * diff
                rec['compliant_dollars_by_premise'][prem] = cs * auth_da
            else:
                rec['overage_by_premise'][prem] = 0
                rec['compliant_dollars_by_premise'][prem] = cs * bb['da']
        
        rec['overage_total'] = sum(rec['overage_by_premise'].values())
        audit.append(rec)
    
    return audit


def summarize_audit(audit):
    """Return totals dict from an audit result list."""
    billed = sum(a['bb']['tot'] for a in audit)
    overage = sum(a['overage_total'] for a in audit)
    return {
        'lines': len(audit),
        'billed': billed,
        'compliant': billed - overage,
        'overage': overage,
        'red_flags': [a for a in audit if a['overage_total'] > 0.5],
        'compliant_lines': [a for a in audit if a['overage_total'] <= 0.5],
    }
